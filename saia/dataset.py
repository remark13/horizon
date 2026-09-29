"""Импорт экспертного Excel-реестра в воспроизводимый нейтральный формат.

Исходный файл не объявляется обучающей выборкой: в нём нет отрицательных
примеров и независимой целевой метки. Импорт сохраняет исходные формулировки,
разбирает ссылки и добавляет только техническую нормализацию.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")


@dataclass(frozen=True)
class EvidenceReference:
    title: str
    url: str
    domain: str
    source_type: str
    language: str | None = None
    published_at: str | None = None
    trust: str | None = None
    observed_or_planned: str = "unknown"


@dataclass(frozen=True)
class SignalSeed:
    source_row: int
    source_number: int
    title: str
    area: str
    companies: str
    expert_rationale: str
    stage_original: str
    stage_code: str
    mention_trend_original: str
    expert_score: float
    evidence: tuple[EvidenceReference, ...]
    label: str | None = None
    label_as_of: str | None = None
    signal_family_id: str | None = None


def stage_code(value: str) -> str:
    text = value.casefold()
    if "исслед" in text:
        return "research"
    if "прототип" in text or "poc" in text:
        return "prototype"
    if "пилот" in text:
        return "pilot"
    if "раннее внедрение" in text:
        return "early_adoption"
    if "масштаб" in text:
        return "scaling"
    return "unknown"


def source_type(domain: str) -> str:
    domain = domain.casefold()
    if domain.endswith("arxiv.org"):
        return "preprint"
    if "openalex.org" in domain:
        return "bibliographic_index"
    if any(token in domain for token in ("nature.com", "science.org", "ieee.org", "acm.org", "springer.com")):
        return "scholarly_publication"
    if any(token in domain for token in ("gov", "europa.eu", "oecd.org", "iso.org")):
        return "official_or_standard"
    if any(token in domain for token in ("techcrunch.com", "siliconangle.com", "venturebeat.com")):
        return "industry_media"
    return "web"


def parse_evidence(value: object) -> tuple[EvidenceReference, ...]:
    text = "" if value is None else str(value)
    found = []
    for title, url in LINK_RE.findall(text):
        clean_url = url.strip().rstrip(".,;")
        domain = urlparse(clean_url).netloc.casefold().removeprefix("www.")
        found.append(EvidenceReference(
            title=title.strip(), url=clean_url, domain=domain,
            source_type=source_type(domain),
        ))
    return tuple(found)


def _header_map(sheet) -> tuple[int, dict[str, int]]:
    aliases = {
        "number": "№",
        "title": "Технология (слабый сигнал)",
        "area": "Область",
        "companies": "Компании",
        "rationale": "Почему это слабый сигнал",
        "stage": "Стадия развития",
        "trend": "Тренд упоминаний",
        "score": "Балл (стадия+тренд)",
        "sources": "Источники",
    }
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        values = {str(value).strip(): idx for idx, value in enumerate(row) if value is not None}
        if aliases["title"] in values:
            missing = [name for name in aliases.values() if name not in values]
            if missing:
                raise ValueError(f"в строке заголовка отсутствуют столбцы: {missing}")
            return row_number, {key: values[name] for key, name in aliases.items()}
    raise ValueError("не найдена строка заголовка датасета")


def load_registry(path: str | Path) -> list[SignalSeed]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    header_row, columns = _header_map(sheet)
    records = []
    for row_number, row in enumerate(
        sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1
    ):
        if row[columns["number"]] is None:
            continue
        records.append(SignalSeed(
            source_row=row_number,
            source_number=int(row[columns["number"]]),
            title=str(row[columns["title"]]).strip(),
            area=str(row[columns["area"]]).strip(),
            companies=str(row[columns["companies"]] or "").strip(),
            expert_rationale=str(row[columns["rationale"]] or "").strip(),
            stage_original=str(row[columns["stage"]] or "").strip(),
            stage_code=stage_code(str(row[columns["stage"]] or "")),
            mention_trend_original=str(row[columns["trend"]] or "").strip(),
            expert_score=float(row[columns["score"]]),
            evidence=parse_evidence(row[columns["sources"]]),
        ))
    return records


def registry_report(path: str | Path, records: list[SignalSeed]) -> dict:
    raw = Path(path).read_bytes()
    domains = Counter(ref.domain for record in records for ref in record.evidence)
    source_types = Counter(ref.source_type for record in records for ref in record.evidence)
    return {
        "file_sha256": hashlib.sha256(raw).hexdigest(),
        "records": len(records),
        "labels_present": sum(record.label is not None for record in records),
        "negative_labels": sum(record.label not in (None, "weak_signal") for record in records),
        "evidence_links": sum(len(record.evidence) for record in records),
        "unique_urls": len({ref.url for record in records for ref in record.evidence}),
        "areas": dict(Counter(record.area for record in records)),
        "scores": dict(Counter(str(record.expert_score) for record in records)),
        "source_types": dict(source_types),
        "top_domains": domains.most_common(20),
        "training_readiness": "not_ready",
        "blocking_reasons": [
            "target label отсутствует",
            "отрицательные и пограничные классы отсутствуют",
            "label_as_of и временные снимки отсутствуют",
            "expert_rationale и expert_score нельзя подавать в признаки из-за target leakage",
        ],
    }


def export_jsonl(records: list[SignalSeed], output: str | Path) -> None:
    destination = Path(output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", encoding="utf-8") as stream:
        for record in records:
            stream.write(json.dumps(asdict(record), ensure_ascii=False) + "\n")


def main() -> int:
    parser = argparse.ArgumentParser(description="Импорт экспертного реестра Horizon")
    parser.add_argument("xlsx")
    parser.add_argument("--output", help="JSONL с нормализованными seed-кейсами")
    parser.add_argument("--report", help="JSON-отчёт о пригодности датасета")
    args = parser.parse_args()
    records = load_registry(args.xlsx)
    report = registry_report(args.xlsx, records)
    if args.output:
        export_jsonl(records, args.output)
    rendered = json.dumps(report, ensure_ascii=False, indent=2)
    if args.report:
        destination = Path(args.report)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text(rendered + "\n", encoding="utf-8")
    print(rendered)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
