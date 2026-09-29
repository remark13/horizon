"""Import the supplied JRC weak-signal workbook as a silver reference set."""
from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from saia.hybrid import digest
from saia.source_registry import file_sha256


VERSION = "jrc-silver-reference-0.4.13"
SHEET = "Слабые сигналы JRC 2024"
HEADERS = (
    "№", "Источник", "Год/издание", "Кластер (рус)",
    "Кластер (англ, оригинал)", "Сигнал (рус)",
    "Сигнал (англ, оригинал)", "Описание (рус)",
    "Есть описание в источнике?", "Ссылка на отчёт",
)


def import_jrc(path: Path, official_claimed_signals: int = 221) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if SHEET not in workbook.sheetnames:
        raise ValueError(f"Нет листа {SHEET!r}.")
    rows = list(workbook[SHEET].iter_rows(values_only=True))
    if not rows or tuple(rows[0]) != HEADERS:
        raise ValueError("Колонки JRC workbook не соответствуют ожидаемому контракту.")
    records = []
    for row_number, row in enumerate(rows[1:], start=2):
        if len(row) != len(HEADERS) or any(row[index] in (None, "") for index in (0, 1, 2, 3, 4, 5, 6, 8, 9)):
            raise ValueError(f"Неполная обязательная строка JRC: {row_number}.")
        description_present = str(row[8]).strip().casefold() == "да"
        description = str(row[7]).strip() if row[7] not in (None, "") else None
        if description_present != (description is not None):
            raise ValueError(f"Противоречие description flag в строке {row_number}.")
        records.append({
            "source_row": row_number,
            "source_number": int(row[0]),
            "source": str(row[1]).strip(),
            "report_year_or_edition": str(row[2]).strip(),
            "cluster_ru": str(row[3]).strip(),
            "cluster_original": str(row[4]).strip(),
            "signal_ru": str(row[5]).strip(),
            "signal_original": str(row[6]).strip(),
            "description_ru": description,
            "description_present_in_supplied_source": description_present,
            "report_url": str(row[9]).strip(),
            "reference_label": "silver_reference_not_outcome_truth",
        })
    originals = Counter(item["signal_original"] for item in records)
    duplicates = [
        {"signal_original": title,
         "source_numbers": [item["source_number"] for item in records
                            if item["signal_original"] == title]}
        for title, count in sorted(originals.items()) if count > 1
    ]
    cluster_counts = Counter(item["cluster_original"] for item in records)
    result = {
        "version": VERSION,
        "source": {
            "file_name": path.name,
            "file_bytes_sha256": file_sha256(path),
            "worksheet": SHEET,
            "report_year": 2024,
            "official_report_url": records[0]["report_url"] if records else None,
            "official_claimed_signals": official_claimed_signals,
        },
        "counts": {
            "rows": len(records),
            "unique_original_titles": len(originals),
            "clusters": len(cluster_counts),
            "descriptions_present": sum(item["description_present_in_supplied_source"] for item in records),
            "descriptions_missing": sum(not item["description_present_in_supplied_source"] for item in records),
            "gap_vs_official_claim": official_claimed_signals - len(records),
        },
        "cluster_counts": dict(sorted(cluster_counts.items())),
        "duplicate_original_titles": duplicates,
        "records": records,
        "allowed_use": ["retrieval_coverage", "cluster_coverage", "manual_reference"],
        "prohibited_use": ["future_success_ground_truth", "automatic_training_labels", "production_score_change"],
        "limitations": [
            "The supplied workbook contains fewer rows than the official report claims.",
            "Missing descriptions remain null and are not generated.",
            "A repeated title may be a cross-cluster entry and is not silently deduplicated.",
            "Report inclusion is not evidence of later scientific or market success.",
        ],
    }
    result["report_payload_sha256"] = digest(result)
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    if args.output.exists():
        parser.error("Reference catalog immutable: use a new versioned output path.")
    result = import_jrc(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("x", encoding="utf-8") as handle:
        json.dump(result, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    print(json.dumps({"output": str(args.output), "counts": result["counts"],
                      "sha256": result["report_payload_sha256"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
