"""Versioned decisions about evidence sources used or considered by SAIA.

The registry is deliberately separate from connector code and scientific
scoring.  A listed source is not automatically trusted, licensed, available,
or suitable for historical evaluation.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path

import yaml
from openpyxl import load_workbook

from saia.hybrid import digest


VERSION = "source-adoption-registry-0.4.51"
SUPPORTED_VERSIONS = {
    "source-adoption-registry-0.4.13",
    "source-adoption-registry-0.4.30",
    "source-adoption-registry-0.4.31",
    "source-adoption-registry-0.4.38",
    "source-adoption-registry-0.4.39",
    "source-adoption-registry-0.4.40",
    "source-adoption-registry-0.4.41",
    "source-adoption-registry-0.4.42",
    "source-adoption-registry-0.4.43",
    "source-adoption-registry-0.4.45",
    "source-adoption-registry-0.4.46",
    "source-adoption-registry-0.4.47",
    "source-adoption-registry-0.4.48",
    "source-adoption-registry-0.4.49",
    "source-adoption-registry-0.4.50",
    VERSION,
}
DECISIONS = {"accept_now", "adapt", "defer", "reject"}
ROLES = {
    "scientific_corpus", "external_attention", "silver_reference",
    "description_only", "future_patent_corpus", "future_software_signal",
    "research_funding_only", "software_diffusion_only",
    "bibliographic_enrichment_only",
    "research_programme_only", "ai_artifact_diffusion_only",
    "news_attention_only", "patent_landscape_only",
    "clinical_translation_only",
    "public_procurement_only",
    "research_artifact_only",
    "investment_lead_only", "company_landscape_only",
}
REQUIRED_FIELDS = {
    "id", "name", "decision", "data", "temporal_coverage",
    "regional_coverage", "access", "credentials", "cost",
    "license_status", "reproducibility", "retrotest_fit", "role", "reason",
}


def file_sha256(path: Path) -> str:
    hasher = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            hasher.update(block)
    return hasher.hexdigest()


def load_registry(path: Path) -> dict:
    value = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict) or value.get("version") not in SUPPORTED_VERSIONS:
        raise ValueError("Нужна поддерживаемая точная версия реестра источников.")
    lineage = [{"file": path.name, "bytes_sha256": file_sha256(path)}]
    if value.get("extends"):
        base_name = value["extends"]
        if not isinstance(base_name, str) or Path(base_name).name != base_name:
            raise ValueError("Базовый реестр должен быть файлом из той же папки.")
        base_path = path.parent / base_name
        base = load_registry(base_path)
        value = dict(value)
        value["sources"] = list(base["sources"]) + list(value.get("sources") or [])
        lineage = list(base.get("_registry_lineage") or []) + lineage
    if value.get("news_publisher_manifest"):
        from saia.news_publishers import registry_rows
        name = value.pop("news_publisher_manifest")
        if name != "news-publishers.v1.json":
            raise ValueError("Неверный файл перечня издателей.")
        manifest = path.parent / name
        value = {**value, "sources": list(value["sources"]) + registry_rows(manifest)}
        lineage.append({"file": name, "bytes_sha256": file_sha256(manifest)})
    replacements = value.get("source_id_replacements") or {}
    if not isinstance(replacements, dict) or any(
        not isinstance(old, str) or not isinstance(new, str) or not old or not new
        for old, new in replacements.items()
    ):
        raise ValueError("Замены ID источников должны быть непустым словарём строк.")
    if replacements:
        value = dict(value)
        value["sources"] = [
            {**source, "id": replacements.get(source.get("id"), source.get("id"))}
            for source in value.get("sources") or []
        ]
    # Re-evaluate a previously considered source without rewriting its history
    # or introducing a second source with the same identity.
    overrides = value.get("source_overrides", {})
    if not isinstance(overrides, dict):
        raise ValueError("Изменения решений по источникам должны быть словарём.")
    known = {source.get("id") for source in value.get("sources") or []}
    for identifier, changes in overrides.items():
        if not isinstance(identifier, str) or identifier not in known:
            raise ValueError("Можно пересмотреть только существующий ID источника.")
        if not isinstance(changes, dict) or not changes or "id" in changes:
            raise ValueError("Пересмотр не должен менять идентичность источника.")
    if overrides:
        value = dict(value)
        value["sources"] = [
            {**source, **overrides.get(source["id"], {})}
            for source in value.get("sources") or []
        ]
    value["_registry_lineage"] = lineage
    if value.get("scientific_score_uses_external_attention") is not False:
        raise ValueError("Внешнее внимание нельзя неявно включать в scientific score.")
    sources = value.get("sources")
    if not isinstance(sources, list) or not sources:
        raise ValueError("Реестр источников пуст.")
    identifiers = []
    for source in sources:
        missing = sorted(REQUIRED_FIELDS - set(source))
        if missing:
            raise ValueError(f"Источник не содержит обязательные поля: {missing}")
        if source["decision"] not in DECISIONS:
            raise ValueError(f"Неизвестное решение по источнику {source['id']}.")
        if source["role"] not in ROLES:
            raise ValueError(f"Неизвестная роль источника {source['id']}.")
        if source["license_status"] == "unknown" and source["decision"] == "accept_now":
            raise ValueError("Источник с неизвестными правами нельзя принять сейчас.")
        identifiers.append(source["id"])
    if len(identifiers) != len(set(identifiers)):
        raise ValueError("ID источников должны быть уникальны.")
    return value


def validate_runtime_sources(registry: dict, runtime_policy: dict) -> dict:
    """Prove that every executable external connector passed the registry gate."""
    approved = {source["id"]: source for source in registry.get("sources") or []}
    configured = {}
    for channel, value in runtime_policy.items():
        if channel in {"version", "scientific_score_modified"}:
            continue
        if not isinstance(value, dict) or not value.get("source"):
            raise ValueError(f"Канал {channel} не содержит канонический source ID.")
        configured[channel] = value["source"]
    missing = sorted(set(configured.values()) - set(approved))
    if missing:
        raise ValueError(
            "Исполняемые источники отсутствуют в adoption registry: " + ", ".join(missing)
        )
    blocked = sorted(
        source_id for source_id in set(configured.values())
        if approved[source_id]["decision"] not in {"accept_now", "adapt"}
    )
    if blocked:
        raise ValueError(
            "Исполняемые источники не допущены к адаптации или использованию: "
            + ", ".join(blocked)
        )
    return {
        "registry_version": registry["version"],
        "runtime_policy_version": runtime_policy.get("version"),
        "configured_channels": configured,
        "configured_source_ids": sorted(set(configured.values())),
        "all_runtime_sources_registered": True,
        "scientific_score_modified": runtime_policy.get("scientific_score_modified"),
    }


def summarize_example_workbook(path: Path) -> dict:
    """Inventory Claude's domain list without treating it as source approval."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    if workbook.sheetnames != ["Источники"]:
        raise ValueError("Неожиданная структура реестра примеров.")
    worksheet = workbook["Источники"]
    rows = list(worksheet.iter_rows(values_only=True))
    expected = (
        "Домен", "Упоминаний", "Категория", "Что это",
        "Доступ к данным (API/скачивание)", "Достоверность",
        "Пример: сигнал", "Пример: ссылка",
    )
    if not rows or tuple(rows[0]) != expected:
        raise ValueError("Неожиданные колонки реестра примеров.")
    records = rows[1:]
    domains = [str(row[0]).strip() for row in records if row[0]]
    if len(domains) != len(set(domains)):
        raise ValueError("В реестре примеров найдены повторяющиеся домены.")
    categories = Counter(str(row[2]).strip() for row in records if row[2])
    mentions = sum(int(row[1]) for row in records if isinstance(row[1], (int, float)))
    return {
        "file_name": path.name,
        "file_bytes_sha256": file_sha256(path),
        "worksheet": "Источники",
        "rows": len(records),
        "unique_domains": len(domains),
        "reported_mentions": mentions,
        "category_counts": dict(sorted(categories.items())),
        "approval_inferred": False,
        "use": "discovery_queue_and_description_examples_only",
        "limitations": [
            "Наличие домена в таблице не подтверждает достоверность публикации.",
            "Права, API, архивная глубина и стабильность требуют отдельной проверки.",
            "Пример ссылки не превращается в наблюдение временного ряда.",
        ],
    }


def build_report(registry: dict, registry_path: Path, examples: dict | None = None) -> dict:
    sources = registry["sources"]
    report = {
        "version": VERSION,
        "reviewed_at": registry["reviewed_at"],
        "registry_file": registry_path.name,
        "registry_file_bytes_sha256": file_sha256(registry_path),
        "registry_lineage": registry.get("_registry_lineage", []),
        "counts": {
            "sources": len(sources),
            "by_decision": dict(sorted(Counter(s["decision"] for s in sources).items())),
            "by_role": dict(sorted(Counter(s["role"] for s in sources).items())),
        },
        "sources": sources,
        "example_source_inventory": examples,
        "scientific_score_uses_external_attention": False,
        "production_thresholds_modified": False,
    }
    report["report_payload_sha256"] = digest(report)
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--registry", type=Path, required=True)
    parser.add_argument("--example-workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    if args.output.exists():
        parser.error("Нужен новый версионированный output; существующий файл не перезаписывается.")
    registry = load_registry(args.registry)
    examples = summarize_example_workbook(args.example_workbook) if args.example_workbook else None
    report = build_report(registry, args.registry, examples)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("x", encoding="utf-8") as handle:
        json.dump(report, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
    print(json.dumps({"output": str(args.output), "counts": report["counts"],
                      "sha256": report["report_payload_sha256"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
