"""Build the SAIA review handoff document from its canonical Markdown."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs/SAIA_ревью_подхода_алгоритма_и_программы_2026-09-24.md"
DEST = SOURCE.with_suffix(".docx")


def hyperlink(paragraph, label: str, href: str) -> None:
    part = paragraph.part
    rid = part.relate_to(href, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "165595")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


TOKEN = re.compile(r"\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|`([^`]+)`|\*([^*]+)\*")


def add_inline(paragraph, content: str) -> None:
    pos = 0
    for match in TOKEN.finditer(content):
        if match.start() > pos:
            paragraph.add_run(content[pos : match.start()])
        if match.group(1) is not None:
            target = match.group(2)
            if not target.startswith("http"):
                target = (SOURCE.parent / target).resolve().as_uri()
            hyperlink(paragraph, match.group(1), target)
        else:
            chunk = next(group for group in match.groups()[2:] if group is not None)
            run = paragraph.add_run(chunk)
            if match.group(3) is not None:
                run.bold = True
            elif match.group(4) is not None:
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            else:
                run.italic = True
        pos = match.end()
    if pos < len(content):
        paragraph.add_run(content[pos:])


def set_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:color"), "D9D9D9")
        borders.append(elem)


def shade_cell(cell, color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shd)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.83)
section.right_margin = Inches(0.83)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.1)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(5.5)
normal.paragraph_format.line_spacing = 1.11

for name, size, before, after in (
    ("Title", 18, 0, 16),
    ("Heading 1", 13, 14, 7),
    ("Heading 2", 11.4, 10, 5),
):
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    if name == "Title":
        ppr = style.element.get_or_add_pPr()
        for border in ppr.findall(qn("w:pBdr")):
            ppr.remove(border)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
idx = 0
while idx < len(lines):
    line = lines[idx].strip()
    if not line:
        idx += 1
        continue
    if line.startswith("| "):
        table_lines = []
        while idx < len(lines) and lines[idx].strip().startswith("|"):
            table_lines.append(lines[idx].strip())
            idx += 1
        rows = [[item.strip() for item in row.strip("|").split("|")] for row in table_lines]
        rows = [row for row in rows if not all(re.fullmatch(r":?-{2,}:?", cell) for cell in row)]
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.autofit = False
        widths = [0.85, 2.50, 3.45] if len(rows[0]) == 3 else [6.8 / len(rows[0])] * len(rows[0])
        for rno, row in enumerate(rows):
            for cno, content in enumerate(row):
                cell = table.cell(rno, cno)
                cell.width = Inches(widths[cno])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.08
                add_inline(paragraph, content)
                for run in paragraph.runs:
                    run.font.size = Pt(9.3)
                    if rno == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                set_cell_border(cell)
                if rno == 0:
                    shade_cell(cell, "193D69")
                elif rno % 2 == 0:
                    shade_cell(cell, "F4F7FA")
        for cell in table.rows[0].cells:
            tr_pr = cell._tc.getparent().get_or_add_trPr()
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(3)
        spacer.paragraph_format.line_spacing = 0.5
        spacer.add_run(" ").font.size = Pt(3)
        continue
    if line.startswith("# "):
        paragraph = doc.add_paragraph(style="Title")
        ppr = paragraph._p.get_or_add_pPr()
        for border in ppr.findall(qn("w:pBdr")):
            ppr.remove(border)
        add_inline(paragraph, line[2:].replace(" — ", " ").replace(",", ""))
    elif line.startswith("## "):
        paragraph = doc.add_paragraph(style="Heading 1")
        if line.startswith("## Очередность следующей работы"):
            paragraph.paragraph_format.page_break_before = True
        add_inline(paragraph, line[3:].replace(" — ", " ").replace(":", ""))
    elif line.startswith("### "):
        paragraph = doc.add_paragraph(style="Heading 2")
        add_inline(paragraph, line[4:].replace(". ", " ").replace(" — ", " ").replace("«", "").replace("»", ""))
    elif line.startswith("- "):
        paragraph = doc.add_paragraph(style="List Bullet")
        add_inline(paragraph, line[2:])
    elif re.match(r"^\d+\. ", line):
        number, content = re.match(r"^(\d+)\. (.*)", line).groups()
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        paragraph.add_run(number + ".  ")
        add_inline(paragraph, content)
    else:
        paragraph = doc.add_paragraph()
        add_inline(paragraph, line)
    idx += 1

doc.save(DEST)
print(DEST)
