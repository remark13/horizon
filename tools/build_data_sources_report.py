from pathlib import Path
import re
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'outputs/SAIA_дополнительные_источники_данных_2026-09-16.md'
OUT = ROOT / 'outputs/SAIA_источники_данных_и_их_ценность.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.75)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(24)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 2'].font.size = Pt(14)
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[name].paragraph_format.keep_with_next = True
doc.core_properties.title = 'SAIA: источники данных и их практическая ценность'
doc.core_properties.subject = 'Свод дополнительных корпусов и размеченных наборов для проверки MVP'
doc.core_properties.author = 'SAIA / Codex'

def inline(p, text):
    pos = 0
    for m in re.finditer(r'\[([^\]]+)\]\(([^\s]+)\)', text):
        p.add_run(text[pos:m.start()])
        h = OxmlElement('w:hyperlink')
        h.set(qn('r:id'), p.part.relate_to(m.group(2), RT.HYPERLINK, is_external=True))
        r = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '0563C1'); props.append(color)
        u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); props.append(u)
        r.append(props)
        t = OxmlElement('w:t'); t.text = m.group(1); r.append(t)
        h.append(r); p._p.append(h)
        pos = m.end()
    p.add_run(text[pos:])
    return p

def p(text, style=None):
    return inline(doc.add_paragraph(style=style), text)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers): cell.text = text
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, text in zip(t.add_row().cells, row): cell.text = text
    for i, row in enumerate(t.rows):
        tr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); tr.append(no_split)
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            pr = cell._tc.get_or_add_tcPr()
            sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), '17365D' if i == 0 else ('F2F5F8' if i % 2 else 'FFFFFF')); pr.append(sh)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge); e.set(qn('w:w'), '85'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                for r in para.runs:
                    r.font.size = Pt(10)
                    if i == 0: r.bold = True; r.font.color.rgb = RGBColor(255,255,255)
    return t

doc.add_paragraph('SAIA: источники данных\nи их практическая ценность', 'Title')
p('Свод для выбора датасетов, подготовки тестов и развития MVP', 'Subtitle')
p('Основание: реестр дополнительного поиска, проверенный 16 сентября 2026 года. Документ не означает, что все наборы скачаны, импортированы или испытаны в SAIA.')
doc.add_heading('Главный вывод', 1)
p('Полноценного универсального датасета «слабый сигнал / шум» среди этих источников не подтверждено. Однако найдены наборы, с помощью которых можно отдельно проверить важные части системы: правильно ли она объединяет статьи, отличает новую идею от пересказа, замечает изменение темы и обосновывает вывод ссылками.')
p('Основной публикационный поток MVP остаётся OpenAlex + arXiv. Внешние наборы нужны прежде всего как испытательные стенды и исторические корпуса. Они не должны подменять основной источник данных или автоматически считаться доказательством будущего рынка.')
p('Первый следующий тест — TRENDNERT, если удастся связать размеченные записи с текстами и датами. Параллельно полезен SciCo для проверки границ кластеров. Для воспроизведения обычного конвейера на исторических AI-публикациях — SciEvo. PreScience перспективен для современного AI-теста, но сначала необходимо найти тематические метки, описанные в статье.')
doc.add_heading('Какая ценность нужна системе', 1)
table(['Задача SAIA', 'Подходящие источники', 'Практический эффект'], [
    ['Различать рост, спад и стабильность', 'TRENDNERT; затем PreScience', 'Проверить, не объявляется ли любая тема растущим трендом'],
    ['Не смешивать разные технологии', 'SciCo / SciCo-Radar', 'Точнее собрать публикации в карточку одного сигнала'],
    ['Отличать новую идею от переформулировки', 'RINoBench; SenticNet', 'Снизить ложную новизну в объяснениях'],
    ['Воспроизвести исторический поиск', 'SciEvo', 'Проверить весь путь от аннотаций до кандидатов'],
    ['Проверять сочетания и обоснования', 'Science4Cast; ITO; SCINLP', 'Развить графовый анализ и доказательность'],
], [2.15, 1.8, 3.05])

doc.add_page_break()
doc.add_heading('1. Свод источников и приоритетов', 1)
p('Приоритеты ниже — рекомендация для нашего MVP, а не оценка общего научного качества источников. «Доступен» означает подтверждение указанной карточки или файла, а не завершённый импорт.')
table(['Источник', 'Что представляет собой', 'Готовность / следующий шаг'], [
 ['TRENDNERT', 'Темы и документы: рост, спад, стабильность', 'Первый CSV скачан. Проверить соответствие хешей текстам и датам'],
 ['SciCo', 'Экспертные связи научных понятий', 'Проверить файлы и контексты; тест синонимов и иерархии'],
 ['SciEvo', 'Большой исторический корпус arXiv', 'Около 8,95 GB. Начать с AI-подвыборки'],
 ['PreScience', 'Современные статьи и будущие траектории', 'Около 963 MB. Тематический файл пока не найден'],
 ['RINoBench', '1 381 идея с оценкой новизны', 'Уточнить лицензию; отдельный тест новизны'],
 ['SenticNet', 'Статьи и синтетические неновые идеи', 'Проверить JSON; разделять семейства статей между выборками'],
 ['Science4Cast', 'Исторический граф AI-понятий', 'Подтвердить архив и лицензию; безопасно обработать pickle'],
 ['ITO', 'Задачи, тесты, результаты и их динамика', 'Проверить исторический массив результатов и условия'],
 ['SCINLP', 'Код проверки научных утверждений', 'Готовый финальный датасет не подтверждён'],
 ['WFtopic', 'Китайские ряды и графы, не AI', 'Резерв; календарь, значения и лицензия требуют проверки'],
], [1.15, 2.45, 3.4])
p('Дополнительные инфраструктурные источники: S2ORC — тексты и ссылки; SciRepEval — проверка эмбеддингов и поиска; Mearman/OpenAlex — стороннее зеркало метаданных. Это не новые эталоны слабых сигналов.')

value = {
 1: ('Главная ценность для SAIA', 'Даёт реальные контрпримеры растущим темам: спад и стабильность. Проверка поможет понять, ранжирует ли система изменения тем, а не просто частотные слова. Первый результат работы — таблица совпадений найденных кластеров с эталонными темами и список причин ошибок.'),
 2: ('Главная ценность для SAIA', 'Наиболее близок к реальному входу MVP: названия и аннотации arXiv. Позволяет повторять один и тот же исторический эксперимент и проверять устойчивость кластеров. Сам по себе не позволяет посчитать точность определения слабых сигналов: эталонные исходы придётся определить отдельно.'),
 3: ('Главная ценность для SAIA', 'Современная AI-среда полезнее старых корпусов для проверки нынешнего языка технологий. Можно сравнивать изменение доли тем и появление их сочетаний. Но пока тематическая разметка не найдена, это кандидат для такого теста, а не готовый эталон.'),
 4: ('Главная ценность для SAIA', 'Проверяет конкретное объяснение в карточке: чем идея отличается от ранее опубликованного. Если модуль плохо проходит этот тест, его показатель новизны нельзя использовать как самостоятельный аргумент в пользу сигнала.'),
 5: ('Главная ценность для SAIA', 'Создаёт трудные отрицательные примеры для семантического анализа: формулировка новая, содержание почти прежнее. Результат теста — доля таких пересказов, ошибочно признанных новыми, и перечень типичных ошибок. Это не оценка прогноза трендов.'),
 6: ('Главная ценность для SAIA', 'Устраняет две практические ошибки: одну технологию показывают несколькими карточками из-за синонимов; разные технологии склеивают из-за общей широкой темы. Качество здесь измеряется связями между понятиями, а не количеством публикаций.'),
 7: ('Главная ценность для SAIA', 'Позволяет испытать гипотезу, что слабый сигнал иногда выражен новым сочетанием известных идей. Проверяем появление конкретной связи в будущем окне. Для редких положительных исходов одного ROC-AUC мало: стоит дополнительно считать precision@K и PR-AUC.'),
 8: ('Главная ценность для SAIA', 'Дополнительная проверка того, превращается ли тема в используемые методы и экспериментальные задачи. Помогает объяснять расхождение: публикаций много, но улучшение конкретного теста остановилось. Не позволяет заключить, что рынок вырос или технология умерла.'),
 9: ('Главная ценность для SAIA', 'Помогает проверять достоверность текста карточки, а не силу сигнала. Например, система пишет «метод снижает стоимость» — проверка должна найти соответствующее утверждение и его ограничения в статье или отметить отсутствие подтверждения.'),
 10: ('Главная ценность для SAIA', 'Небольшие реальные графы удобны для дешёвого инженерного теста алгоритмов. Однако экономические и политические темы нельзя использовать как подтверждение качества обнаружения AI-сигналов или как полноценное покрытие азиатских AI-исследований.'),
}

text = SOURCE.read_text()
parts = re.split(r'^## ', text, flags=re.M)
for part in parts:
    match = re.match(r'(\d+)\. (.+)\n', part)
    if not match: continue
    num = int(match.group(1))
    title = match.group(2)
    doc.add_page_break()
    doc.add_heading(f'2.{num}. {title}', 1)
    doc.add_heading('Практическая ценность ' + title.split(':')[0], 2)
    p(value[num][1])
    body = part[match.end():].strip()
    # Markdown paragraphs retain the research record and inline source links.
    for block in re.split(r'\n\s*\n', body):
        lines = block.splitlines()
        if lines[0].startswith('- '):
            for line in lines:
                p(line.removeprefix('- '), 'List Bullet')
        else:
            p(' '.join(line.strip() for line in lines))

doc.add_page_break()
doc.add_heading('3. Вспомогательные источники', 1)
aux = next(part for part in parts if part.startswith('Дополнительные источники самого'))
for block in re.split(r'\n\s*\n', aux.split('\n',1)[1].strip()):
    p(' '.join(block.splitlines()))
doc.add_heading('Что они дают — и чего не дают', 2)
p('S2ORC помогает восстановить текст и контекст ссылок; SciRepEval — выбрать более подходящее научное представление текста; зеркало OpenAlex — упростить пакетное чтение. Ни один из них не создаёт автоматически независимую разметку «перспективный слабый сигнал». Для каждого надо отдельно подтвердить доступ, версию и условия выбранных данных.')

doc.add_page_break()
doc.add_heading('4. Как превратить этот свод в полезные тесты', 1)
doc.add_heading('4.1. Сначала проверить доступность и соответствие входу MVP', 2)
p('Для каждого выбранного набора составить паспорт: URL, версия или коммит, лицензия данных отдельно от лицензии кода, контрольные суммы, доступные поля, дата публикации и дата снимка. Затем взять небольшую выборку и проверить, можно ли восстановить названия, аннотации, даты и идентификаторы OpenAlex/arXiv. Посчитать долю успешно связанных записей и показать несвязанные отдельно.')
doc.add_heading('4.2. Разделить проверяемые задачи', 2)
p('Не обучать одну модель на всех перечисленных метках. Для роста использовать тематические траектории; для семантики — новизну и связи понятий; для текста карточки — подтверждение утверждений. У каждого теста должны быть своя единица оценки и отдельный результат.')
table(['Что оцениваем', 'Рекомендуемые показатели', 'Как читать результат'], [
 ['Рост / спад / стабильность', 'MAP; precision@K; ошибки по классам', 'Насколько полезен верх списка; какие устойчивые темы система ошибочно считает растущими'],
 ['Качество кластеров', 'Парные precision / recall / F1; ручной разбор', 'Склейка разных технологий и дробление синонимов'],
 ['Новизна', 'Согласие со шкалой; ложная новизна', 'Отличает ли модуль содержание от новой формулировки'],
 ['Новые графовые связи', 'ROC-AUC вместе с PR-AUC и precision@K', 'Доля полезных прогнозов при редких положительных исходах'],
 ['Доказательность карточки', 'Доля подтверждённых утверждений; ошибки ссылок', 'Есть ли в источнике заявленное преимущество и ограничение'],
], [1.45,2.2,3.35])
p('Показатели в этой таблице — предложение для разработки. Они ещё не измерены на перечисленных наборах и не являются заявленными результатами SAIA.')
doc.add_heading('4.3. Проверить раннее обнаружение без подсказки из будущего', 2)
p('Выбрать историческую дату проверки. Вход системы ограничить публикациями и версиями, которые уже существовали к этой дате. Цитирования, документы и результаты более позднего периода использовать только для оценки последствий. Не давать системе сегодняшнюю аннотацию обновлённого препринта или сегодняшнее число цитат как будто они были известны раньше.')
p('На этапе прототипа разумно сравнить SAIA с простыми правилами: ростом числа публикаций, ростом доли темы и случайным ранжированием. Проверять не только успешные истории, но и стабильные и снижающиеся темы. Фиксировать время раннего предупреждения и причины ложных срабатываний. Если исторические признаки восстановить нельзя, назвать эксперимент ретроспективным, а не прогнозным.')
doc.add_page_break()
doc.add_heading('4.4. Не считать все отрицательные примеры шумом', 2)
p('Хранить раздельно: тема стабильна; тема падает; не достигла будущего порога; идея уже существовала; утверждение не подтверждено; контрпример синтетический; исход неизвестен. Стабильная технология может быть полезной и иметь большой рынок. Неновая идея может быстро распространиться. Поэтому отрицательная метка одной задачи не равна бесполезному сигналу во всех задачах.')
doc.add_heading('Рекомендуемый порядок', 2)
p('1. TRENDNERT: доступ и связывание, затем тест роста, спада и стабильности. 2. SciCo: качество объединения понятий. 3. SciEvo: историческая AI-подвыборка и повторяемый конвейер. 4. PreScience: поиск тематических меток, затем современный AI-тест. 5. RINoBench и SenticNet: новизна. 6. Science4Cast, ITO и SCINLP: графы и доказательность после базовых проверок.')
p('Итоговая ценность этого набора источников — не обещание автоматически предсказать будущие технологии, а возможность последовательно выявлять слабые места MVP и подтверждать улучшения отдельными воспроизводимыми тестами.')

footer = sec.footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run('SAIA • источники данных | '); run.font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
# Keep headings plain and descriptive; canonical dataset names remain in prose.
for para in doc.paragraphs:
    if para.style.name in ['Title', 'Subtitle'] or para.style.name.startswith('Heading'):
        para.text = re.sub(r'[^\w\s]', ' ', para.text)
        para.text = re.sub(r' +', ' ', para.text).strip()
doc.save(OUT)
# The bundled Word base can carry a title-rule residue. Remove paragraph borders
# after serialization, without altering explicit table cell borders.
with ZipFile(OUT) as z:
    entries = {name: z.read(name) for name in z.namelist()}
for name in ['word/document.xml', 'word/styles.xml']:
    root = etree.fromstring(entries[name])
    for border in root.xpath('//w:pBdr', namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}):
        border.getparent().remove(border)
    entries[name] = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
with ZipFile(OUT, 'w') as z:
    for name, content in entries.items(): z.writestr(name, content)
with ZipFile(OUT) as z:
    assert z.testzip() is None
print(OUT)
print('Hyperlinks:', sum(len(p._p.xpath('.//w:hyperlink')) for p in doc.paragraphs))
