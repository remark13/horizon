from __future__ import annotations

import json
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/a/Documents/Codex/2026-09-09/ruen-15-ml/saia_codex_v0_3")
OUT_DIR = ROOT / "outputs/01a08680-0825-74d2-8b5f-65c057bbdceb"
OUT = OUT_DIR / "Результаты_ретротеста_8_сигналов_SAIA.docx"
CHART = OUT_DIR / "retrotest_funnel.png"

NAVY = "17365D"
MID_BLUE = "5B9BD5"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "F4F7FB"
PALE_GRAY = "F2F2F2"
GRID = "D9D9D9"
GREEN = "70AD47"
AMBER = "FFC000"
RED = "C00000"

CASES = [
    {
        "id": "registry-073-mobile-personalization", "n": 73,
        "name": "Дообучение и multi LoRA на смартфоне",
        "status": "Поиск не восстановил контрольные работы",
        "result": "Не пройден",
        "analysis": (
            "Problem-first запрос по on-device AI и mobile LLM сформировал полезный корпус, "
            "но ни одна из двух контрольных публикаций не попала в выдачу OpenAlex. Поэтому "
            "кластеризатор не имел возможности обнаружить целевую линию. Это ошибка полноты "
            "поиска, а не доказательство отсутствия сигнала."
        ),
        "action": (
            "Расширить поиск семантическими соседями для embedded AI и смартфонов, не включая "
            "слова fine-tuning и LoRA. Затем повторить тест на том же срезе."
        ),
        "sources": [
            ("2512.08211", "MobileFineTuner", "09.12.2025", "https://arxiv.org/abs/2512.08211"),
            ("2604.18655", "Multi LoRA edge deployment", "20.04.2026", "https://arxiv.org/abs/2604.18655"),
        ],
    },
    {
        "id": "registry-076-memory-defense", "n": 76,
        "name": "Защита от отравления памяти агентов",
        "status": "Поиск не восстановил контрольные работы и корпус слишком мал",
        "result": "Не пройден",
        "analysis": (
            "Обе статьи существуют и опубликованы до даты среза, но problem-first запрос их не "
            "вернул. После фильтра качества осталось 39 работ и только три тематические линии, "
            "что ниже минимальной базы сравнения. Найденная слабая линия по agentic AI security "
            "не может считаться подтверждением memory poisoning без контрольных публикаций."
        ),
        "action": (
            "Добавить отдельное расширение запроса по persistent memory и retrieval agents, "
            "а полноту поиска измерять до кластеризации."
        ),
        "sources": [
            ("2606.12703", "Certified Defence Against Runtime Memory Poisoning", "10.06.2026", "https://arxiv.org/abs/2606.12703"),
            ("2605.22842", "When Memory Poisoning Looks Like Model Failure", "12.05.2026", "https://arxiv.org/abs/2605.22842"),
        ],
    },
    {
        "id": "registry-078-plc-codegen", "n": 78,
        "name": "Генерация кода ПЛК на IEC 61131 3",
        "status": "Обе работы найдены и соединены но тема слишком широкая",
        "result": "Не пройден",
        "analysis": (
            "Система нашла обе контрольные статьи, допустила их после проверки качества и "
            "поместила в один воспроизводимый кластер. Однако в этом кластере 41 работа, поэтому "
            "контрольные статьи составили 4,88 процента при пороге 10 процентов. Кластер описывает "
            "PLC в целом, а не самостоятельную линию LLM генерации Structured Text."
        ),
        "action": (
            "Добавить генератор методических кандидатов по устойчивым фразам, совместной "
            "встречаемости LLM, compiler feedback и Structured Text внутри широкого PLC кластера."
        ),
        "sources": [
            ("2410.22159", "Training LLMs for IEC 61131 3 Structured Text", "29.10.2024", "https://arxiv.org/abs/2410.22159"),
            ("2412.02410", "AutoPLC", "03.12.2024", "https://arxiv.org/abs/2412.02410"),
        ],
    },
    {
        "id": "registry-081-industrial-rl", "n": 81,
        "name": "Обучение управления на реальном промышленном объекте",
        "status": "Одна работа найдена но осталась шумом",
        "result": "Не пройден",
        "analysis": (
            "OpenAlex вернул статью об управлении промышленной энергосистемой, но она не вошла "
            "ни в один устойчивый кластер. Журнальная статья о self-learning control не попала в "
            "problem-first выдачу. Корпус дал только три тематические линии, поэтому сравнительные "
            "метрики недостаточны."
        ),
        "action": (
            "Расширить предметный корпус публикациями об industrial robots и live control, "
            "сохранив запрет на прямой запрос reinforcement learning и PPO."
        ),
        "sources": [
            ("2605.31044", "RL for Controlling Industrial Energy Systems", "29.05.2026", "https://arxiv.org/abs/2605.31044"),
            ("10.1038/s41598-025-28904-8", "Software-defined self-learning control", "13.01.2026", "https://doi.org/10.1038/s41598-025-28904-8"),
        ],
    },
    {
        "id": "registry-092-trading-benchmarks", "n": 92,
        "name": "Бенчмарки LLM агентов для трейдинга",
        "status": "Контрольные работы не восстановлены поиском",
        "result": "Не пройден",
        "analysis": (
            "После фильтра качества осталось 416 публикаций и 33 тематические линии, то есть "
            "сравнительная база была достаточной. Тем не менее problem-first запрос по algorithmic "
            "trading не вернул ни одну из трёх контрольных работ. Тест завершился на этапе поиска."
        ),
        "action": (
            "Ввести измеряемое расширение от algorithmic trading к autonomous and LLM agents. "
            "Слова benchmark и stress test должны оставаться запрещёнными до оценки результата."
        ),
        "sources": [
            ("2603.00285", "TraderBench", "27.02.2026", "https://arxiv.org/abs/2603.00285"),
            ("2510.11695", "When Agents Trade", "13.10.2025", "https://arxiv.org/abs/2510.11695"),
            ("2605.28359", "Memory-Controlled Benchmark for LLM Trading Agents", "27.05.2026", "https://arxiv.org/abs/2605.28359"),
        ],
    },
    {
        "id": "registry-094-confidential-edge", "n": 94,
        "name": "Confidential computing на слабых edge устройствах",
        "status": "Обе работы образовали самостоятельную линию",
        "result": "Пройден",
        "analysis": (
            "Обе контрольные статьи были найдены, прошли фильтр качества и попали в один "
            "воспроизводимый кластер. Они составили 40 процентов кластера при пороге 10 процентов. "
            "Временная дисциплина, карантин, повторяемость и остальные инфраструктурные проверки "
            "пройдены. Это единственный полный положительный результат из восьми кейсов."
        ),
        "action": (
            "Сохранить кейс как регрессионный положительный контроль. Любая следующая версия "
            "конвейера должна повторно находить эту линию без добавления целевых терминов."
        ),
        "sources": [
            ("2606.07470", "Confidential DNN Inference on Low-End Edge Devices", "05.06.2026", "https://arxiv.org/abs/2606.07470"),
            ("2605.03213", "Confidential Computing for Agentic AI", "04.05.2026", "https://arxiv.org/abs/2605.03213"),
        ],
    },
    {
        "id": "registry-095-in-sensor", "n": 95,
        "name": "In sensor и processing in pixel вычисления",
        "status": "Обе работы найдены но разошлись по разным темам",
        "result": "Не пройден",
        "analysis": (
            "Обе журнальные статьи присутствовали в корпусе и были допущены к моделированию. "
            "Обзор edge intelligence попал в широкий кластер edge computing, а аппаратная работа "
            "про processing-in-pixel — в кластер sensors and neuromorphic sensing. Система увидела "
            "две предметные области, но не связала их в одну сквозную технологическую линию."
        ),
        "action": (
            "Добавить связи по ключевым техническим фразам и цитатному графу поверх предметных "
            "кластеров. Семантических эмбеддингов одной статьи недостаточно для сквозного метода."
        ),
        "sources": [
            ("10.1038/s44335-025-00040-6", "Edge intelligence through in-sensor computing", "01.10.2025", "https://doi.org/10.1038/s44335-025-00040-6"),
            ("10.1038/s44335-025-00029-1", "ADC-less processing-in-pixel", "18.06.2025", "https://doi.org/10.1038/s44335-025-00029-1"),
        ],
    },
    {
        "id": "registry-096-speculative-mobile", "n": 96,
        "name": "Speculative decoding с выгрузкой весов",
        "status": "Поисковый корпус оказался непригодно мал",
        "result": "Не пройден",
        "analysis": (
            "Запрос вернул 16 записей, после фильтра качества и временного среза осталось пять. "
            "Контрольные статьи в выдачу не попали, а единственная тематическая линия не даёт "
            "сравнительной базы. Такой результат нельзя интерпретировать как отсутствие сигнала."
        ),
        "action": (
            "Расширить problem-first поиск до on-device inference и smartphone deployment, "
            "но не использовать speculative decoding, Lever и SLED в запросе."
        ),
        "sources": [
            ("2605.16786", "Lever", "16.05.2026", "https://arxiv.org/abs/2605.16786"),
            ("2506.09397", "SLED", "11.06.2025", "https://arxiv.org/abs/2506.09397"),
        ],
    },
]


def load_report(case: dict) -> dict:
    path = ROOT / "reports/generated" / f"{case['id']}-benchmark.json"
    return json.loads(path.read_text(encoding="utf-8"))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text: str, url: str, color="0563C1"):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(c)
    r_pr.append(u)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_col_widths(table, widths) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def style_table(table, header=True, font_size=8.5) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r_idx, row in enumerate(table.rows):
        set_cant_split(row)
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_idx == 0 and header:
                set_cell_shading(cell, NAVY)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
                    run.font.size = Pt(font_size)
                    if r_idx == 0 and header:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
    if header:
        set_repeat_table_header(table.rows[0])


def add_heading(doc, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_metric_table(doc, rows, left_header="Метрика", right_header="Наблюдение"):
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = left_header
    table.rows[0].cells[1].text = right_header
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)
    set_col_widths(table, [Inches(3.4), Inches(3.2)])
    style_table(table, font_size=9)
    return table


def create_chart() -> None:
    labels = ["Не восстановлены контрольные работы", "Частичное восстановление", "Найдены, но не выделены", "Полный тест пройден"]
    values = [4, 1, 2, 1]
    colors = ["#C00000", "#ED7D31", "#FFC000", "#70AD47"]
    image = Image.new("RGB", (1500, 530), "white")
    draw = ImageDraw.Draw(image)
    font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
    bold_path = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
    font = ImageFont.truetype(font_path, 30)
    bold = ImageFont.truetype(bold_path, 32)
    small = ImageFont.truetype(font_path, 25)
    left, top, bar_height, gap, max_width = 610, 60, 62, 54, 700
    for index, (label, value, color) in enumerate(zip(labels, values, colors)):
        y = top + index * (bar_height + gap)
        draw.text((30, y + 12), label, font=font, fill="#333333")
        width = int(max_width * value / 4)
        draw.rounded_rectangle((left, y, left + width, y + bar_height), radius=9, fill=color)
        draw.text((left + width + 18, y + 10), str(value), font=bold, fill="#222222")
    draw.text((left, 495), "Число кейсов", font=small, fill="#666666")
    image.save(CHART)


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for case in CASES:
        case["report"] = load_report(case)
    create_chart()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(31, 31, 31)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, before, after in (("Title", 20, 0, 12), ("Heading 1", 15, 16, 7), ("Heading 2", 12, 11, 5)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title = doc.add_paragraph(style="Title")
    title.add_run("Результаты ретроспективного теста восьми слабых технологических сигналов")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = subtitle.add_run("SAIA MVP   OpenAlex и arXiv   Проверка 16 сентября 2026 года")
    run.bold = True
    run.font.color.rgb = RGBColor(89, 89, 89)

    p = doc.add_paragraph()
    r = p.add_run("Главный вывод. ")
    r.bold = True
    p.add_run(
        "Полный тест прошёл один кейс из восьми — confidential computing на слабых edge устройствах. "
        "Два кейса дошли до контрольных публикаций, но система не выделила их как самостоятельную "
        "линию. В пяти кейсах результат сорвался раньше из-за неполноты problem-first поиска или "
        "недостаточного корпуса. Реестр полезен как набор проверочных гипотез, но пока не является "
        "готовой gold-разметкой или обучающим датасетом."
    )

    add_heading(doc, "Что именно проверялось", 1)
    doc.add_paragraph(
        "Для каждого кейса создан отдельный исторический срез. В поисковый запрос входила проблема "
        "или предметная область, но не название целевого метода. Контрольные arXiv ID и DOI "
        "использовались только после эмбеддингов, кластеризации и расчёта кандидатов. Успех требовал "
        "не простого присутствия статьи, а совместного попадания контрольных работ в отдельную тему, "
        "достаточную долю этой темы и прохождение временных и качественных проверок."
    )
    chart_shape = doc.add_picture(str(CHART), width=Inches(6.85))
    chart_shape._inline.docPr.set(
        "descr",
        "Воронка ретроспективного теста: восемь кейсов, семь непройденных и один полностью пройденный.",
    )
    chart_shape._inline.docPr.set("title", "Результат ретроспективного теста SAIA")
    caption = doc.paragraphs[-1]
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph("Рисунок 1. Итоги восьми problem-first ретротестов")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)

    add_heading(doc, "Сводные результаты", 1)
    table = doc.add_table(rows=1, cols=6)
    for cell, text in zip(table.rows[0].cells, ["№", "Сигнал", "Корпус", "Контроль", "Темы", "Результат"]):
        cell.text = text
    for case in CASES:
        report = case["report"]
        recovery = report["target_recovery"]
        found = sum(1 for item in report["target_publications"] if item.get("found_in_corpus"))
        cells = table.add_row().cells
        values = [
            case["n"], case["name"],
            f"{report['corpus']['canonical_works']} всего\n{sum(report['corpus']['quality_decisions'].get(k, 0) for k in ('include',))} допущено",
            f"{found}/{len(report['target_publications'])} найдено\n{recovery['max_target_works_in_one_topic']} вместе",
            report["discovery"]["topic_lines"], case["result"],
        ]
        for cell, value in zip(cells, values):
            cell.text = str(value)
        color = GREEN if case["result"] == "Пройден" else RED
        cells[-1].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(color)
        cells[-1].paragraphs[0].runs[0].font.bold = True
    set_col_widths(table, [Inches(0.38), Inches(2.12), Inches(1.05), Inches(1.05), Inches(0.55), Inches(1.05)])
    style_table(table, font_size=8.2)

    p = doc.add_paragraph()
    r = p.add_run("Как читать результат. ")
    r.bold = True
    p.add_run(
        "Ноль найденных контрольных работ означает провал retrieval recall. Две найденные работы в "
        "разных кластерах означают провал связывания. Две работы в одном слишком большом кластере "
        "означают, что система увидела область, но не выделила слабый сигнал."
    )

    add_heading(doc, "Методика теста", 1)
    steps = [
        "Проверены 14 arXiv записей через официальный arXiv API и три журнальные статьи Nature через DOI и OpenAlex.",
        "Для каждого кейса задана дата среза сразу после появления двух контрольных публикаций. Публикации на дату среза и позже исключались.",
        "OpenAlex собирал полный результат по problem-first запросу в заданном периоде. Целевые методы и идентификаторы в запрос не включались.",
        "Дубликаты объединялись по DOI, arXiv ID и заголовку с авторами. Затем применялся фильтр релевантности и карантин дат и источников.",
        "Для допущенных работ построены SPECTER2 proximity эмбеддинги. BERTopic работал по квартальным окнам на диагностическом масштабе от двух работ.",
        "Каждая кластеризация повторена дважды. Состав тем совпал во всех восьми миссиях.",
        "Контрольные ID и DOI были открыты только после расчёта кандидатов. Порог самостоятельности — не менее двух работ в одной теме и не менее 10 процентов состава этой темы.",
    ]
    for step in steps:
        doc.add_paragraph(step, style="List Number")

    p = doc.add_paragraph()
    r = p.add_run("Ограничение теста. ")
    r.bold = True
    p.add_run(
        "Независимым каналом обнаружения в этом прогоне был OpenAlex. arXiv использовался для "
        "проверки первых версий контрольных работ. Закреплённое помесячное зеркало arXiv, применённое "
        "в предыдущем историческом тесте, не содержит публикаций 2026 года. Поэтому этот результат "
        "нельзя представлять как объединённый слепой мониторинг двух полных источников."
    )

    doc.add_page_break()
    add_heading(doc, "Результаты по каждому сигналу", 1)
    for index, case in enumerate(CASES):
        if index:
            doc.add_page_break()
        report = case["report"]
        recovery = report["target_recovery"]
        found = sum(1 for item in report["target_publications"] if item.get("found_in_corpus"))
        add_heading(doc, f"Сигнал {case['n']} {case['name']}", 2)
        p = doc.add_paragraph()
        r = p.add_run(f"{case['result']}. ")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(GREEN if case["result"] == "Пройден" else RED)
        p.add_run(case["status"] + ".")
        doc.add_paragraph(case["analysis"])
        add_metric_table(doc, [
            ("Дата среза", report["as_of_date"]),
            ("Канонических работ", report["corpus"]["canonical_works"]),
            ("Допущено после фильтра качества", report["corpus"]["quality_decisions"].get("include", 0)),
            ("Тематических линий", report["discovery"]["topic_lines"]),
            ("Контрольных работ найдено", f"{found} из {len(report['target_publications'])}"),
            ("Максимум контрольных работ в одной теме", recovery["max_target_works_in_one_topic"]),
            ("Доля контрольных работ в доминирующей теме", "н/д" if recovery["dominant_target_share"] is None else f"{100 * recovery['dominant_target_share']:.2f}%"),
            ("Инфраструктурные проверки", "пройдены" if report["pipeline_controls_passed"] else "не пройдены"),
        ])
        doc.add_paragraph("Контрольные первоисточники")
        for identifier, title, date, url in case["sources"]:
            p = doc.add_paragraph(style="List Bullet")
            add_hyperlink(p, f"{title}   {identifier}", url)
            p.add_run(f"   первая публикация {date}")
        p = doc.add_paragraph()
        r = p.add_run("Следующий проверяемый шаг. ")
        r.bold = True
        p.add_run(case["action"])

    doc.add_page_break()
    add_heading(doc, "Что показал тест для продукта", 1)
    findings = [
        ("Поиск сейчас ограничивает результат сильнее кластеризации", "В четырёх кейсах контрольные работы существовали в OpenAlex, но problem-first запрос не включил их в корпус. Ещё один кейс восстановлен только частично."),
        ("Присутствие статьи не равно обнаружению сигнала", "В кейсе PLC две контрольные работы оказались в одной теме, но составили 4,88 процента широкого кластера. Без порога доли система ошибочно объявила бы успех."),
        ("Сквозные методы распадаются по областям применения", "In-sensor computing разделился между edge computing и аппаратными сенсорами. SPECTER2 и BERTopic хорошо группируют предметные области, но хуже связывают общий метод между ними."),
        ("Малый корпус нельзя компенсировать высоким score", "В кейсах memory defence, industrial RL и speculative mobile оказалось меньше пяти сопоставимых тем. Процентили и ранжирование в такой выборке ненадёжны."),
        ("Положительный контроль существует", "Confidential edge прошёл поиск, качество, кластеризацию, порог доли и повторяемость. Этот кейс следует использовать как обязательный регрессионный тест."),
    ]
    for heading, body in findings:
        p = doc.add_paragraph()
        r = p.add_run(heading + ". ")
        r.bold = True
        p.add_run(body)

    add_heading(doc, "Ошибки прототипа найденные во время теста", 1)
    fixes = [
        ("Идентификатор arXiv внутри DOI", "OpenAlex часто хранит препринт как DOI 10.48550/arXiv. Код не создавал отдельный arXiv ID, поэтому benchmark не находил уже загруженную работу. Парсер исправлен и покрыт тестом."),
        ("BERTopic на единственном кластере", "min_df равный двум делал раннее микроокно математически недопустимым. Для кластерных документов установлен min_df равный единице."),
        ("UMAP на двух и трёх работах", "Спектральная инициализация падала на малых окнах. Размерность теперь адаптируется, а окна из двух и трёх работ используют детерминированный граф сходства с порогом из паспорта методики."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Ошибка"
    table.rows[0].cells[1].text = "Исправление"
    for name, fix in fixes:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = fix
    set_col_widths(table, [Inches(2.1), Inches(4.5)])
    style_table(table, font_size=9)
    doc.add_paragraph("После исправлений локальный набор проверок содержит 78 успешных тестов и один пропущенный интеграционный тест без базы данных.")

    add_heading(doc, "Приоритеты следующей версии", 1)
    priorities = [
        ("P0", "Измерять полноту поиска отдельно", "Добавить retrieval recall по контрольным публикациям, журнал расширения запроса и причину невключения каждой контрольной работы."),
        ("P0", "Подключить независимый актуальный arXiv контур", "Использовать официальный bulk или OAI источник с датой первой версии. OpenAlex enrichment по уже найденным ID не считать независимым подтверждением."),
        ("P0", "Расширять problem-first запрос без утечки ответа", "Генерировать соседние проблемы, объекты и ограничения. Целевые методы и названия контрольных работ хранить в закрытом benchmark блоке."),
        ("P1", "Добавить гибридный генератор кандидатов", "Соединить устойчивые фразы, граф совместной встречаемости, цитатные связи и семантические кластеры. Это необходимо для PLC и in-sensor кейсов."),
        ("P1", "Добавить отрицательные и пограничные кейсы", "Для каждого положительного сигнала подобрать сопоставимую тему, которая не выросла. Иначе Precision at 15 и частота ложных тревог не измеряются."),
    ]
    table = doc.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Приоритет", "Задача", "Проверяемый результат"]):
        cell.text = text
    for priority, task, outcome in priorities:
        cells = table.add_row().cells
        cells[0].text = priority
        cells[1].text = task
        cells[2].text = outcome
    set_col_widths(table, [Inches(0.75), Inches(2.2), Inches(3.75)])
    style_table(table, font_size=8.8)

    add_heading(doc, "Критерии следующего прогона", 1)
    criteria = [
        "Problem-first retrieval восстанавливает обе контрольные публикации минимум в шести из восьми кейсов.",
        "Кейс №94 остаётся положительным регрессионным контролем без добавления целевых терминов.",
        "№78 выделяется из общего PLC кластера или получает честный статус candidate с объяснением недостаточной самостоятельности.",
        "№95 связывает аппаратную и обзорную работу через гибридный граф либо явно показывает две независимые линии.",
        "Во всех миссиях отсутствуют публикации после даты среза, карантин не попадает в темы, а повторный запуск даёт тот же состав кластеров.",
        "Отдельный набор отрицательных кейсов позволяет рассчитать Precision at 15 и долю ложных тревог."
    ]
    for item in criteria:
        doc.add_paragraph(item, style="List Bullet")

    add_heading(doc, "Итоговое решение", 1)
    doc.add_paragraph(
        "Восемь выбранных сигналов следует сохранить как регрессионный набор SAIA. Текущую версию "
        "нельзя оценивать одной цифрой accuracy: четыре кейса проверяют полноту поиска, два — "
        "способность выделять сквозной метод, один — работу на малом корпусе и один — полный "
        "положительный сценарий. Главная разработческая задача следующего цикла — повысить recall "
        "problem-first сбора и добавить гибридное выделение методических линий, не ослабляя as-of "
        "дисциплину и требования к самостоятельности темы."
    )

    doc.add_page_break()
    add_heading(doc, "Техническая воспроизводимость", 1)
    add_metric_table(doc, [
        ("Дата выполнения", "16.09.2026"),
        ("Источник корпуса", "OpenAlex API"),
        ("Проверка контрольных версий", "официальный arXiv API и DOI"),
        ("Эмбеддинги", "SPECTER2 proximity 3447645e и 20815596"),
        ("Кластеризация", "BERTopic по кварталам scale seed"),
        ("Повторяемость", "два одинаковых прогона каждой миссии"),
        ("Методика", "SAIA v0.1 hash 56cc17631d40730d"),
        ("Локальные тесты", "78 passed 1 skipped"),
    ], left_header="Параметр", right_header="Зафиксированное значение")

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("SAIA MVP   Ретроспективный тест восьми сигналов   16.09.2026")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(117, 117, 117)

    core = doc.core_properties
    core.title = "Результаты ретроспективного теста восьми слабых технологических сигналов"
    core.subject = "SAIA MVP OpenAlex arXiv"
    core.author = "SAIA project team"
    core.keywords = "SAIA, weak signals, OpenAlex, arXiv, retrospective benchmark"

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
